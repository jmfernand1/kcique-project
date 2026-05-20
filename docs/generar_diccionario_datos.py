# -*- coding: utf-8 -*-
"""
Generador del Diccionario de Datos de Integración (KCIQUE).

Produce el archivo `Diccionario_Datos_Integracion.xlsx`, que indica a los
equipos externos QUÉ INFORMACIÓN se requiere para registrar Desembolsos,
Cargos Fijos, Garantías (app dexter) y Casos de Débito (app adagio), insumo
de las automatizaciones que se basan en lo ingestado en la base de datos.

El libro contiene, por cada entidad:
  - Una hoja "Dicc - <Entidad>"      -> diccionario de campos.
  - Una hoja "Plantilla - <Entidad>" -> formato para que el equipo lo llene.
Más una hoja "Instrucciones" de portada.

Uso:
    python generar_diccionario_datos.py

Regenerar este archivo cada vez que cambien los modelos/serializers.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# El .xlsx se genera siempre junto a este script (carpeta docs/).
ARCHIVO_SALIDA = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Diccionario_Datos_Integracion.xlsx",
)

# Columnas del diccionario
COLUMNAS = ["Campo", "Descripción", "Tipo de dato", "Obligatorio",
            "Formato / Regla", "Valores permitidos", "Ejemplo"]

# ---------------------------------------------------------------------------
# Definición de entidades. Cada campo:
# (campo, descripcion, tipo, obligatorio, formato, valores_permitidos, ejemplo)
# ---------------------------------------------------------------------------

DESEMBOLSO = {
    "titulo": "DESEMBOLSO",
    "endpoint": "POST /dexter/api/desembolsos/",
    "nota": ("Crea el desembolso base. Los cargos fijos se registran aparte "
             "(ver hoja CargoFijo) referenciando el id de este desembolso."),
    "campos": [
        ("referencia", "Identificador único del desembolso.", "Texto (máx 100)", "Sí", "Único, no se puede repetir", "", "DES-001"),
        ("obligacion", "Número de obligación / crédito.", "Entero", "Sí", "Solo dígitos", "", "9001234567"),
        ("id_cliente", "Identificador del cliente.", "Entero", "Sí", "Solo dígitos", "", "1234567"),
        ("nit_beneficiario", "NIT del beneficiario del desembolso.", "Entero", "Sí", "Solo dígitos, sin dígito de verificación", "", "8001234561"),
        ("aliado", "Nombre del aliado comercial.", "Texto (máx 255)", "Sí", "", "", "Aliado XYZ"),
        ("tipo_cta_destino", "Tipo de cuenta destino.", "Texto (máx 50)", "Sí", "", "AHORROS / CORRIENTE", "AHORROS"),
        ("cod_tipo_cuenta_destino", "Código del tipo de cuenta destino.", "Entero", "Sí", "", "", "1"),
        ("num_cta_destino", "Número de cuenta destino.", "Entero", "No", "Solo dígitos", "", "123456789"),
        ("banco_destino", "Nombre del banco destino.", "Texto (máx 100)", "No", "", "", "Bancolombia"),
        ("cod_banco_destino", "Código del banco destino.", "Entero", "No", "", "", "7"),
        ("valor_desembolso", "Valor a desembolsar.", "Decimal", "Sí", "Mayor a 0, punto como separador decimal", "", "15000000.0"),
        ("numero_tramos", "Número de tramos del desembolso.", "Entero", "Sí", "", "", "1"),
        ("dia_pago_cuota", "Día del mes de pago de la cuota.", "Entero", "No", "Entre 1 y 31", "", "5"),
        ("estado", "Estado del desembolso.", "Texto (máx 50)", "No", "Lo gestiona el sistema si se omite", "", "PENDIENTE"),
        ("tipo_desembolso", "Id del tipo de desembolso.", "Entero (referencia)", "No", "Debe existir en Tipos de Desembolso", "", "1"),
        ("observaciones", "Observaciones libres.", "JSON / objeto", "No", "Objeto JSON", "", '{"nota": "texto"}'),
        ("ids_sharepoint", "Identificadores de documentos en SharePoint.", "JSON / objeto", "No", "Objeto JSON", "", '{"doc1": "url"}'),
        ("etapa_desembolso", "Etapa interna del desembolso.", "Entero", "No", "Por defecto 0", "", "0"),
        ("intentos_desembolso", "Número de intentos de desembolso.", "Entero", "No", "Por defecto 0", "", "0"),
    ],
}

CARGO_FIJO = {
    "titulo": "CARGO FIJO",
    "endpoint": "POST /dexter/api/cargos-fijos/",
    "nota": ("Cada cargo fijo se asocia a un desembolso ya creado mediante su id. "
             "Las fechas usan formato numérico YYYYMMDD."),
    "campos": [
        ("desembolso", "Id del desembolso al que pertenece el cargo.", "Entero (referencia)", "Sí", "Debe existir un Desembolso con ese id", "", "10"),
        ("codigo", "Código del cargo fijo.", "Texto (máx 10)", "Sí", "", "", "CF01"),
        ("nombre_cargo_fijo", "Nombre o descripción del cargo.", "Texto (máx 255)", "Sí", "", "", "Cuota de manejo"),
        ("fecha_efectiva", "Fecha en que aplica el cargo.", "Fecha", "No", "Formato YYYYMMDD (también acepta YYYY-MM-DD)", "", "20260131"),
        ("fecha_revision", "Fecha de revisión del cargo.", "Fecha", "No", "Formato YYYYMMDD (también acepta YYYY-MM-DD)", "", "20260630"),
        ("periodicidad", "Periodicidad de cobro del cargo.", "Texto (máx 50)", "No", "", "MENSUAL / TRIMESTRAL / ANUAL", "MENSUAL"),
        ("valor", "Valor del cargo fijo.", "Decimal", "No", "Punto como separador decimal", "", "12000.0"),
    ],
}

GARANTIA = {
    "titulo": "GARANTÍA",
    "endpoint": "POST /dexter/api/garantias/",
    "nota": "Garantía vehicular asociada a una obligación. Las fechas usan formato YYYYMMDD.",
    "campos": [
        ("referencia", "Identificador único de la garantía.", "Texto (máx 100)", "Sí", "Único, no se puede repetir", "", "GAR-001"),
        ("obligacion", "Número de obligación asociada.", "Entero", "Sí", "Solo dígitos", "", "9001234567"),
        ("id_cliente", "Identificador del cliente.", "Entero", "Sí", "Solo dígitos", "", "1234567"),
        ("id_garante", "Identificador del garante.", "Entero", "Sí", "Solo dígitos", "", "7654321"),
        ("cod_fasecolda", "Código Fasecolda del vehículo.", "Texto (máx 50)", "Sí", "", "", "08042099"),
        ("codigo_fasecolda", "Código Fasecolda (campo secundario).", "Texto (máx 50)", "Sí", "", "", "08042099"),
        ("placa", "Placa del vehículo.", "Texto (máx 20)", "Sí", "", "", "ABC123"),
        ("color", "Color del vehículo.", "Texto (máx 100)", "No", "", "", "Blanco"),
        ("poliza_vehiculo", "Número de póliza del vehículo.", "Texto (máx 50)", "No", "", "", "POL-998"),
        ("cod_aseguradora", "Código de la aseguradora.", "Texto (máx 20)", "No", "", "", "ASE01"),
        ("nro_poliza", "Número de póliza.", "Texto (máx 50)", "No", "", "", "NP-12345"),
        ("valor_asegurado", "Valor asegurado del vehículo.", "Decimal", "No", "", "", "50000000.0"),
        ("fecha_vencimiento_seguro", "Fecha de vencimiento del seguro.", "Fecha", "No", "Formato YYYYMMDD", "", "20270131"),
        ("tipo_prima", "Tipo de prima.", "Texto (máx 50)", "No", "", "", "ANUAL"),
        ("valor_prima", "Valor de la prima.", "Decimal", "No", "", "", "1200000.0"),
        ("folio_electronico", "Folio electrónico.", "Texto (máx 50)", "No", "", "", "FE-0001"),
        ("valor_vehiculo", "Valor comercial del vehículo.", "Entero", "No", "", "", "60000000"),
        ("modelo", "Modelo (año) del vehículo.", "Texto (máx 10)", "No", "", "", "2024"),
        ("fecha_prenda", "Fecha de la prenda.", "Fecha", "No", "Formato YYYYMMDD", "", "20260201"),
        ("chasis", "Número de chasis.", "Texto (máx 50)", "No", "", "", "CH123456"),
        ("motor", "Número de motor.", "Texto (máx 50)", "No", "", "", "MT123456"),
        ("serie", "Número de serie.", "Texto (máx 50)", "No", "", "", "SR123456"),
        ("servicio", "Tipo de servicio del vehículo.", "Texto (máx 50)", "No", "", "PARTICULAR / PÚBLICO", "PARTICULAR"),
        ("fecha_desembolso", "Fecha del desembolso asociado.", "Fecha", "No", "Formato YYYYMMDD", "", "20260131"),
        ("estado", "Estado de la garantía.", "Texto (máx 50)", "No", "", "", "PENDIENTE"),
    ],
}

CASO_DEBITO = {
    "titulo": "CASO DE DÉBITO (adagio)",
    "endpoint": "POST /adagio/api/casos/",
    "nota": ("Caso de débito que alimenta las automatizaciones de adagio. Solo "
             "cod_caso_bizagi es obligatorio para la API, pero el negocio "
             "requiere los datos de cuenta/banco para poder procesar el caso. "
             "Los campos de trazabilidad (fechas, intentos, errores) los llena "
             "el sistema y no deben enviarse."),
    "campos": [
        ("cod_caso_bizagi", "Código único del caso en Bizagi.", "Texto (máx 255)", "Sí", "Único, no se puede repetir", "", "CD-2026-0001"),
        ("num_prestamo", "Número de préstamo asociado al caso.", "Texto (máx 255)", "No", "Requerido por negocio", "", "9001234567"),
        ("docsoldv", "Documentos solicitados o validados.", "Texto (máx 255)", "No", "", "", "DOC-001"),
        ("doctitulardv", "Número de identificación del titular de la cuenta.", "Texto (máx 255)", "No", "Requerido por negocio", "", "1234567"),
        ("tipo_de_cuenta", "Tipo de cuenta.", "Texto (máx 100)", "No", "Requerido por negocio", "AHORROS / CORRIENTE", "AHORROS"),
        ("numcta_debito", "Número de cuenta para débito.", "Texto (máx 255)", "No", "Requerido por negocio", "", "123456789"),
        ("secuencia_cta", "Secuencia de la cuenta.", "Texto (máx 100)", "No", "", "", "01"),
        ("codigo_del_banco", "Código del banco.", "Texto (máx 100)", "No", "Requerido por negocio", "", "007"),
        ("codigo_ciudad", "Código de la ciudad.", "Texto (máx 100)", "No", "", "", "11001"),
        ("tipo_debito", "Tipo de débito.", "Texto (lista cerrada)", "No", "Por defecto AL TITULAR", "AL TITULAR / A TERCEROS / NO APLICA", "AL TITULAR"),
        ("autoriza", "Indica quién autoriza el débito.", "Texto (máx 10)", "No", "", "", "SI"),
        ("fecha_desembolso", "Fecha en que se realizó el desembolso.", "Texto (máx 10)", "No", "Texto de fecha", "", "2026-01-31"),
        ("estado", "Estado inicial del caso.", "Texto (lista cerrada)", "No", "Por defecto PENDIENTE", "PENDIENTE / PENDIENTE DATOS OK / GRABADO / PENDIENTE BIZAGI / FINALIZADO / VALIDAR / CON ERROR", "PENDIENTE"),
        ("datos_adicionales", "Información adicional en formato JSON.", "JSON / objeto", "No", "Objeto JSON", "", '{"clave": "valor"}'),
    ],
}

ENTIDADES = [DESEMBOLSO, CARGO_FIJO, GARANTIA, CASO_DEBITO]

# Listas desplegables para las plantillas (campo -> opciones)
LISTAS = {
    "tipo_debito": ["AL TITULAR", "A TERCEROS", "NO APLICA"],
    "estado_caso": ["PENDIENTE", "PENDIENTE DATOS OK", "GRABADO",
                    "PENDIENTE BIZAGI", "FINALIZADO", "VALIDAR", "CON ERROR"],
}

# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
AZUL = "1F4E78"
AZUL_CLARO = "D9E2F3"
AMARILLO = "FFF2CC"
GRIS = "F2F2F2"
VERDE = "E2EFDA"

f_titulo = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
f_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
f_normal = Font(name="Calibri", size=10)
f_oblig = Font(name="Calibri", size=10, bold=True, color="C00000")
f_nota = Font(name="Calibri", size=9, italic=True, color="404040")

fill_titulo = PatternFill("solid", fgColor=AZUL)
fill_header = PatternFill("solid", fgColor=AZUL)
fill_oblig = PatternFill("solid", fgColor=AMARILLO)
fill_plantilla = PatternFill("solid", fgColor=VERDE)

borde = Border(*[Side(style="thin", color="BFBFBF")] * 4)
wrap = Alignment(wrap_text=True, vertical="top")
centro = Alignment(horizontal="center", vertical="center")


def estilo_celda(celda, fuente=f_normal, alineacion=wrap, relleno=None):
    celda.font = fuente
    celda.alignment = alineacion
    celda.border = borde
    if relleno:
        celda.fill = relleno


# ---------------------------------------------------------------------------
# Construcción del libro
# ---------------------------------------------------------------------------
def crear_hoja_instrucciones(wb):
    ws = wb.active
    ws.title = "Instrucciones"
    ws.sheet_properties.tabColor = AZUL
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110

    filas = [
        ("Diccionario de Datos de Integración - KCIQUE", "titulo"),
        ("", None),
        ("¿Para qué sirve este archivo?", "sub"),
        ("Indica QUÉ INFORMACIÓN se requiere para registrar la data que alimenta "
         "las automatizaciones de KCIQUE (apps dexter y adagio).", "txt"),
        ("", None),
        ("¿Cómo está organizado?", "sub"),
        ("Por cada entidad hay dos hojas:", "txt"),
        ("  •  'Dicc - <Entidad>': diccionario de campos (descripción, tipo, "
         "si es obligatorio, formato, valores permitidos y ejemplo).", "txt"),
        ("  •  'Plantilla - <Entidad>': formato para llenar los datos. La fila 2 "
         "trae un ejemplo; bórrela y registre sus datos desde la fila 2.", "txt"),
        ("", None),
        ("Entidades incluidas", "sub"),
        ("  •  Desembolso  ->  POST /dexter/api/desembolsos/", "txt"),
        ("  •  Cargo Fijo  ->  POST /dexter/api/cargos-fijos/", "txt"),
        ("  •  Garantía    ->  POST /dexter/api/garantias/", "txt"),
        ("  •  Caso de Débito (adagio)  ->  POST /adagio/api/casos/", "txt"),
        ("", None),
        ("Reglas importantes", "sub"),
        ("  •  Los campos marcados como Obligatorio = Sí no pueden ir vacíos.", "txt"),
        ("  •  Los campos 'referencia' y 'cod_caso_bizagi' son ÚNICOS: no se "
         "pueden repetir entre registros.", "txt"),
        ("  •  Las fechas de dexter (fecha_efectiva, fecha_prenda, etc.) usan "
         "formato numérico YYYYMMDD. Ejemplo: 31 de enero de 2026 = 20260131.", "txt"),
        ("  •  Los campos con 'Valores permitidos' solo aceptan esos valores.", "txt"),
        ("  •  No diligencie campos de trazabilidad (fechas de proceso, intentos, "
         "errores): los completa el sistema automáticamente.", "txt"),
        ("", None),
        ("Cómo entregar la información", "sub"),
        ("Llene las hojas 'Plantilla - <Entidad>' y devuelva este archivo al "
         "equipo de integración. El detalle técnico de la API está en "
         "dexter/API_DOCUMENTATION.md y en el esquema OpenAPI /api/schema/.", "txt"),
    ]

    fila = 1
    for texto, tipo in filas:
        c = ws.cell(row=fila, column=2, value=texto)
        if tipo == "titulo":
            c.font = f_titulo
            c.fill = fill_titulo
            c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
            ws.row_dimensions[fila].height = 28
        elif tipo == "sub":
            c.font = Font(name="Calibri", size=11, bold=True, color=AZUL)
        elif tipo == "txt":
            c.font = f_normal
            c.alignment = Alignment(wrap_text=True, vertical="top")
        fila += 1


def crear_hoja_diccionario(wb, entidad):
    nombre = "Dicc - " + entidad["titulo"].split(" (")[0].title()
    ws = wb.create_sheet(title=nombre[:31])
    anchos = [26, 46, 20, 13, 34, 40, 22]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    t = ws.cell(row=1, column=1, value=entidad["titulo"])
    t.font = f_titulo
    t.fill = fill_titulo
    t.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 26

    # Endpoint + nota
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    e = ws.cell(row=2, column=1, value="Endpoint API:  " + entidad["endpoint"])
    e.font = Font(name="Calibri", size=10, bold=True, color=AZUL)
    e.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    e.alignment = Alignment(vertical="center", indent=1)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
    n = ws.cell(row=3, column=1, value="Nota:  " + entidad["nota"])
    n.font = f_nota
    n.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[3].height = 42

    # Encabezados
    fila_header = 4
    for col, titulo in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=fila_header, column=col, value=titulo)
        estilo_celda(c, fuente=f_header, alineacion=centro, relleno=fill_header)
    ws.row_dimensions[fila_header].height = 22

    # Filas de campos
    fila = fila_header + 1
    for campo in entidad["campos"]:
        relleno_fila = fill_plantilla if campo[3] == "Sí" else None
        for col, valor in enumerate(campo, start=1):
            c = ws.cell(row=fila, column=col, value=valor)
            if col == 4:  # Obligatorio
                if valor == "Sí":
                    estilo_celda(c, fuente=f_oblig, alineacion=centro, relleno=fill_oblig)
                else:
                    estilo_celda(c, alineacion=centro)
            else:
                estilo_celda(c, relleno=relleno_fila)
        fila += 1

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{fila_header}:G{fila - 1}"


def crear_hoja_plantilla(wb, entidad):
    nombre = "Plantilla - " + entidad["titulo"].split(" (")[0].title()
    ws = wb.create_sheet(title=nombre[:31])
    campos = entidad["campos"]

    for i, campo in enumerate(campos, start=1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = max(16, len(campo[0]) + 4)
        # Encabezado = nombre del campo
        h = ws.cell(row=1, column=i, value=campo[0])
        estilo_celda(h, fuente=f_header, alineacion=centro,
                     relleno=fill_header if campo[3] != "Sí" else PatternFill("solid", fgColor="C00000"))
        # Fila de ejemplo
        ej = ws.cell(row=2, column=i, value=campo[6])
        estilo_celda(ej, fuente=Font(name="Calibri", size=10, italic=True, color="808080"))

    # Nota de la fila de ejemplo
    ws.cell(row=4, column=1,
            value="^ Fila 2 = EJEMPLO. Bórrela y registre sus datos desde la fila 2. "
                  "Las columnas con encabezado rojo son obligatorias.").font = f_nota

    # Validaciones de lista
    if entidad is CASO_DEBITO:
        cols = {c[0]: idx + 1 for idx, c in enumerate(campos)}
        _agregar_validacion(ws, cols.get("tipo_debito"), LISTAS["tipo_debito"])
        _agregar_validacion(ws, cols.get("estado"), LISTAS["estado_caso"])

    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "70AD47"


def _agregar_validacion(ws, col_idx, opciones):
    if not col_idx:
        return
    letra = get_column_letter(col_idx)
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(opciones),
                        allow_blank=True, showDropDown=False)
    dv.error = "Valor no permitido. Seleccione uno de la lista."
    dv.prompt = "Seleccione un valor de la lista."
    ws.add_data_validation(dv)
    dv.add(f"{letra}2:{letra}500")


def main():
    wb = Workbook()
    crear_hoja_instrucciones(wb)
    for entidad in ENTIDADES:
        crear_hoja_diccionario(wb, entidad)
    for entidad in ENTIDADES:
        crear_hoja_plantilla(wb, entidad)
    wb.save(ARCHIVO_SALIDA)
    print("Archivo generado:", ARCHIVO_SALIDA)
    print("Hojas:", ", ".join(ws.title for ws in wb.worksheets))


if __name__ == "__main__":
    main()
